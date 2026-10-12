import hashlib
import io
import json
import tempfile
import unittest
import warnings
from datetime import datetime
from itertools import permutations
from pathlib import Path

import openpyxl
from pypdf import PdfReader
from PIL import Image as PILImage

from central.core import Store, read_excel, date_value
from central.pdf import generate_pdf
from central.reports import totals, filter_rows, group_rows, leaves, sort_rows

ROOT=Path(__file__).resolve().parents[1]
warnings.filterwarnings('ignore',message='Workbook contains no default style')
HEADERS=['Nº de Pedido da Plataforma','Nº do Subpedido','Nº de Pedido','Plataformas','Nome da Loja no UpSeller',
         'Estado do Pedido','Hora do Pedido','Impressão da Etiqueta','Prazo de Envio','Nome do Anúncio',
         'ID do Anúncio','SKU','Variação','ID da Variante','Qtd. do Produto','Nome do Destinatário']


def excel(records, title='sheet', extra=False):
    w=openpyxl.Workbook();w.active.title=title
    w.active.append(HEADERS)
    for changes in records:
        base={'Nº de Pedido da Plataforma':'00012345678901234567890','Nº do Subpedido':'',
              'Nº de Pedido':'UP123','Plataformas':'Shopee','Nome da Loja no UpSeller':'AnnaLiv',
              'Estado do Pedido':'Retirada','Hora do Pedido':'2026-10-10 10:00:00',
              'Impressão da Etiqueta':'2026-10-12 10:00:00','Prazo de Envio':'2026-10-13 12:00:00',
              'Nome do Anúncio':'Produto <comercial> & nome longo','ID do Anúncio':'AD1',
              'SKU':'100.Preto.G','Variação':'Preto,G','ID da Variante':'VAR1','Qtd. do Produto':1,
              'Nome do Destinatário':'SEGREDO_CLIENTE_NAO_RETER'}
        base.update(changes)
        w.active.append([base.get(h,'') for h in HEADERS])
    if extra: w.create_sheet('vazia')
    b=io.BytesIO();w.save(b);return b.getvalue()


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.store=Store(Path(self.temp.name)/'test.sqlite3')

    def tearDown(self):
        self.store.close();self.temp.cleanup()

    @unittest.skipUnless((ROOT/'amostras/upseller').exists() and (ROOT/'amostras/catalogo/catalogo_base_validada_2026-10-12.csv').exists(), 'Amostras privadas disponíveis somente na máquina da operação')
    def test_real_reference_all_lines_catalog_and_reimport(self):
        files=[(p.name,p.read_bytes()) for p in sorted((ROOT/'amostras/upseller').glob('*.xlsx'))]
        hashes={name:hashlib.sha256(b).hexdigest() for name,b in files}
        result=self.store.import_files(files,'2026-10-12T18:00:00+00:00')
        self.assertTrue(all('error' not in f for f in result['files']),result)
        t=totals(self.store.rows())
        self.assertEqual((t['orders'],t['items'],t['confirmed'],t['pending']),(163,179,180,0))
        self.assertEqual(sum(f['rows'] for f in result['files']),254)
        self.assertEqual(sum(f['duplicate'] for f in result['files']),2)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM observations').fetchone()[0],254)
        for f in result['files']:self.assertEqual(len(self.store.ledger(f['attempt'])),f['rows'])
        self.assertEqual(sum(r['confirmed_at_import'] for f in result['files'] for r in self.store.ledger(f['attempt'])),180)
        self.store.import_files(files,'2026-10-14T18:00:00+00:00')
        self.assertEqual(totals(self.store.rows()),t)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM observations').fetchone()[0],508)
        for row in self.store.rows():
            self.assertEqual(row['first_seen'],'2026-10-12T18:00:00+00:00')
            self.assertEqual(row['last_seen'],'2026-10-14T18:00:00+00:00')
            self.assertIn('Já importado',row['labels'])
        catalog=(ROOT/'amostras/catalogo/catalogo_base_validada_2026-10-12.csv').read_bytes()
        self.assertEqual(self.store.import_catalog(catalog)['accepted'],87)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM products').fetchone()[0],87)
        missing=[r for r in self.store.rows() if r['candidate_ref']=='857']
        self.assertEqual(len(missing),3)
        self.assertTrue(all(any('ausente' in s for s in r['issues']) for r in missing))
        conflict=[r for r in self.store.rows() if r['sku']=='875-Preto-M']
        self.assertTrue(any('Conflito REF 875' in s for s in conflict[0]['issues']))
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM links').fetchone()[0],0)
        self.assertEqual(hashes,{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'amostras/upseller').glob('*.xlsx')})

    def test_multiple_products_same_order_and_quantity_two(self):
        self.store.import_files([('a.xlsx',excel([{}, {'SKU':'100.Branco.G','ID da Variante':'VAR2','Variação':'Branco,G'},
                                                 {'SKU':'101.Vinho.M','ID do Anúncio':'AD2','Qtd. do Produto':2}]))])
        self.assertEqual(totals(self.store.rows())['confirmed'],4)
        self.assertEqual(totals(self.store.rows())['items'],3)
        self.assertEqual(totals(self.store.rows())['orders'],1)

    def test_identical_rows_preserved_pending_and_manual_reconciliation(self):
        self.store.import_files([('a.xlsx',excel([{},{}]))])
        rows=self.store.rows();self.assertEqual(len(rows),2)
        self.assertEqual(totals(rows)['confirmed'],0);self.assertEqual(totals(rows)['pending'],2)
        self.store.resolve(rows[0]['item_id'],'distinct')
        self.store.resolve(rows[1]['item_id'],'duplicate',rows[0]['item_id'])
        self.assertEqual(totals(self.store.rows())['confirmed'],1)
        self.assertEqual(len(self.store.rows()),2)
        self.assertEqual(totals(self.store.rows())['resolved'],1)

    def test_quantity_divergence_requires_decision(self):
        self.store.import_files([('a.xlsx',excel([{}]))])
        self.store.import_files([('b.xlsx',excel([{'Qtd. do Produto':2}]))])
        rows=self.store.rows();self.assertEqual(totals(rows)['confirmed'],1)
        self.assertEqual(totals(rows)['pending'],2)
        pending=next(r for r in rows if r['reconciliation']=='pending')
        self.store.resolve(pending['item_id'],'replace')
        self.assertEqual(totals(self.store.rows())['confirmed'],2)
        self.assertEqual(totals(self.store.rows())['pending'],0)
        self.assertEqual(len(self.store.history(rows[0]['order_id'])),2)
        self.store.import_files([('c.xlsx',excel([{'Qtd. do Produto':2,'Estado do Pedido':'Conferir'}]))])
        self.assertEqual(totals(self.store.rows())['confirmed'],2)
        self.assertEqual(totals(self.store.rows())['pending'],0)

    def test_repeated_pending_observation_does_not_inflate_pending(self):
        self.store.import_files([('a.xlsx',excel([{}]))])
        self.store.import_files([('b.xlsx',excel([{'Qtd. do Produto':2}]))])
        self.store.import_files([('c.xlsx',excel([{'Qtd. do Produto':2,'Estado do Pedido':'Outro'}]))])
        # Há um item confirmado e uma hipótese pendente. Repetir essa hipótese não soma duas vezes.
        self.assertEqual(totals(self.store.rows())['confirmed'],1)
        self.assertEqual(totals(self.store.rows())['pending'],2)
        self.assertEqual(len(self.store.rows()),2)

    def test_reexport_multiset_preserves_pending_multiplicity(self):
        self.store.import_files([('a.xlsx',excel([{},{}]))])
        self.store.import_files([('b.xlsx',excel([{'Estado do Pedido':'Outro'},{'Estado do Pedido':'Outro'}]))])
        self.assertEqual(len(self.store.rows()),2)
        self.assertEqual(totals(self.store.rows())['pending'],2)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM observations').fetchone()[0],4)

    def test_new_signature_existing_order_requires_review(self):
        self.store.import_files([('a.xlsx',excel([{}]))])
        self.store.import_files([('b.xlsx',excel([{'SKU':'101.Branco.M','ID do Anúncio':'OUTRO','ID da Variante':'OUTRA'}]))])
        self.assertEqual(totals(self.store.rows())['confirmed'],1)
        self.assertEqual(totals(self.store.rows())['pending'],1)
        self.assertEqual(len(self.store.rows()),2)

    def test_missing_sku_keeps_visible_pending(self):
        self.store.import_files([('a.xlsx',excel([{'SKU':''}]))])
        self.assertEqual(len(self.store.rows()),1)
        self.assertEqual(totals(self.store.rows())['pending'],1)

    def test_reexport_changed_status_no_new_quantity_and_full_history(self):
        self.store.import_files([('a.xlsx',excel([{'Estado do Pedido':'Em Processo','Impressão da Etiqueta':''}]))],'2026-10-10T18:00:00+00:00')
        self.store.import_files([('b.xlsx',excel([{}]))],'2026-10-12T18:00:00+00:00')
        rows=self.store.rows();self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['status'],'Retirada')
        self.assertEqual(len(self.store.history(rows[0]['order_id'])),2)
        self.assertIn('Etiqueta impressa',rows[0]['labels'])

    def test_partial_export_does_not_remove_items(self):
        self.store.import_files([('a.xlsx',excel([{}, {'SKU':'101.Branco.G','ID do Anúncio':'AD2'}]))])
        self.store.import_files([('b.xlsx',excel([{}],extra=True))])
        self.assertEqual(totals(self.store.rows())['items'],2)
        self.assertEqual(totals(self.store.rows())['confirmed'],2)

    def test_same_external_number_different_shops_and_marketplaces(self):
        self.store.import_files([('a.xlsx',excel([{}, {'Nome da Loja no UpSeller':'DonaMora'}, {'Plataformas':'Kwai'}]))])
        self.assertEqual(totals(self.store.rows())['orders'],3)
        self.assertTrue(all(r['order']=='00012345678901234567890' for r in self.store.rows()))

    def test_incomplete_identity_bad_quantity_unknown_state_no_omission(self):
        self.store.import_files([('a.xlsx',excel([{'Nº de Pedido da Plataforma':'','Qtd. do Produto':'errada','Estado do Pedido':'???','Impressão da Etiqueta':''},
                                                 {'Prazo de Envio':'inválido','Qtd. do Produto':-1}]))])
        rows=self.store.rows();self.assertEqual(len(rows),2)
        self.assertEqual(totals(rows)['invalid_quantity'],2)
        self.assertIn('Status desconhecido',rows[0]['labels'])
        self.assertIn('Impressão não informada',rows[0]['labels'])
        self.assertIsNone(rows[1]['deadline'])

    def test_bad_file_recorded_and_other_files_import(self):
        result=self.store.import_files([('bad.xlsx',b'bad'),('good.xlsx',excel([{}]))])
        self.assertIn('error',result['files'][0]);self.assertEqual(len(self.store.rows()),1)
        self.assertEqual(self.store.db.execute('SELECT COUNT(*) FROM attempts').fetchone()[0],2)

    def test_reordered_headers_missing_and_numeric_ids(self):
        w=openpyxl.load_workbook(io.BytesIO(excel([{}])))
        sheet=w.active
        values=list(sheet.values)
        sheet.delete_rows(1,sheet.max_row)
        for row in values:sheet.append(list(reversed(row)))
        output=io.BytesIO();w.save(output)
        records,_=read_excel(output.getvalue())
        self.assertEqual(records[0]['raw']['sku'],'100.Preto.G')
        self.assertEqual(records[0]['quantity'],1)
        self.store.import_files([('numeric.xlsx',excel([{'Nº de Pedido da Plataforma':1234567890123456}]))])
        self.assertEqual(totals(self.store.rows())['pending'],1)
        self.assertTrue(any('numérico' in issue for issue in self.store.rows()[0]['issues']))

    def test_manual_marks_persist_and_no_automatic_exclusion(self):
        content=excel([{}]);self.store.import_files([('a.xlsx',content)],'2026-10-10T18:00:00+00:00')
        oid=self.store.rows()[0]['order_id']
        for mark in ['Pendente','Em produção','Produzido','Separado','Conferido','Não necessita ação']:
            self.store.set_mark(oid,mark,'2026-10-11T18:00:00+00:00')
            self.store.import_files([('a.xlsx',content)],'2026-10-12T18:00:00+00:00')
            row=self.store.rows('2026-10-14T18:00:00+00:00')[0]
            self.assertEqual(row['mark'],mark)
            self.assertEqual(totals(filter_rows(self.store.rows()))['confirmed'],1)
            self.assertEqual(row['age_days'],4)
            if mark=='Pendente':self.assertEqual(row['operational_days'],3)
            else:self.assertIsNone(row['operational_days'])

    def test_catalog_manual_fields_and_reviewed_links_protected(self):
        content=b'REF;PRODUTO;CUSTO_TOTAL;COR\n100;Nome;10;Inventada\n'
        self.store.import_catalog(content)
        self.store.save_product('100','Revisado','11.50','Body')
        self.store.import_catalog(content)
        product=dict(self.store.db.execute('SELECT * FROM products').fetchone())
        self.assertEqual(product['name'],'Revisado');self.assertEqual(product['cost'],'11.50')
        self.store.import_files([('a.xlsx',excel([{}]))])
        self.assertEqual(self.store.rows()[0]['ref'],'Não validada')
        self.store.save_link('Shopee','AnnaLiv','100.Preto.G','100','Preto','G')
        self.assertEqual(self.store.rows()[0]['ref'],'100')
        self.assertEqual(self.store.rows()[0]['color'],'Preto')
        self.assertEqual(self.store.rows()[0]['category'],'Body')

    def test_personal_fields_not_persisted(self):
        self.store.import_files([('a.xlsx',excel([{}]))])
        dump='\n'.join(self.store.db.iterdump())
        self.assertNotIn('SEGREDO_CLIENTE_NAO_RETER',dump)
        pdf=generate_pdf(self.store.rows())
        self.assertNotIn('SEGREDO_CLIENTE_NAO_RETER',''.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages))

    def test_backup_and_favorite_reopen(self):
        self.store.import_files([('a.xlsx',excel([{}]))]);self.store.save_favorite('Modelo',{'groups':['marketplace','deadline']})
        dest=Path(self.temp.name)/'backup.sqlite3';self.store.backup(dest)
        other=Store(dest)
        try:
            self.assertEqual(totals(other.rows()),totals(self.store.rows()))
            self.assertEqual(json.loads(other.db.execute('SELECT config FROM favorites').fetchone()[0]),{'groups':['marketplace','deadline']})
        finally:other.close()

    def test_grouping_permutations_and_hour_filter(self):
        self.store.import_files([('a.xlsx',excel([{}, {'Nº de Pedido da Plataforma':'2','Prazo de Envio':'2026-10-13 23:59:59'},
                                                 {'Nº de Pedido da Plataforma':'3','Plataformas':'Kwai'}]))])
        rows=self.store.rows()
        for dims in permutations(['marketplace','deadline','color','ref']):
            tree=group_rows(rows,list(dims));flattened=[r for leaf in leaves(tree) for r in leaf['rows']]
            self.assertEqual(totals(flattened),totals(rows));self.assertEqual(len(flattened),3)
        for dims in [['deadline','category','color'],['category','color','ref'],['shop','marketplace','deadline'],['color','ref','size']]:
            self.assertEqual(group_rows(rows,dims)['totals'],totals(rows))
        limited=filter_rows(rows,{'deadline_to':'2026-10-13T12:00:00-03:00'})
        self.assertEqual(len(limited),2)
        self.assertEqual(sort_rows(rows)[-1]['order'],'2')
        self.assertEqual(len(filter_rows(rows,{'marketplace':['Kwai']})),1)
        with self.assertRaises(ValueError):group_rows(rows,['marketplace','marketplace'])

    def test_pdf_a4_multpage_labels_photos_and_totals(self):
        self.store.import_files([('a.xlsx',excel([{'Nº de Pedido da Plataforma':str(i),'Nome do Anúncio':'Descrição longa '*12} for i in range(35)]))])
        photo=Path(self.temp.name)/'foto.png';PILImage.new('RGB',(300,400),'purple').save(photo)
        self.store.save_product('100','Produto foto','10','Body',str(photo))
        self.store.save_link('Shopee','AnnaLiv','100.Preto.G','100','Preto','G')
        for photographs in [False,True]:
            pdf=generate_pdf(self.store.rows(),{'groups':['marketplace','deadline'],'photos':photographs,'labels':True})
            reader=PdfReader(io.BytesIO(pdf));self.assertGreater(len(reader.pages),1)
            content='\n'.join(p.extract_text() for p in reader.pages)
            self.assertIn('TOTAL GERAL: 35 pedidos distintos',content)
            self.assertIn('35 unidades confirmadas',content)
            self.assertIn('Etiqueta impressa',content)
            self.assertIn('13/10/2026 12:00:00',content)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.width),595.2756,places=2)
            self.assertTrue(all('Marketplace / Loja' in p.extract_text() for p in reader.pages[:-1]))
            if photographs:self.assertTrue(any('/XObject' in p['/Resources'] for p in reader.pages))


if __name__=='__main__':unittest.main()
